"""Tracker addendum for session 3: the DMG build run and its one toolchain note."""

import datetime as dt

import openpyxl

NOW = dt.datetime(2026, 9, 28, 21, 30).strftime("%Y-%m-%d %H:%M")
TODAY = "2026-09-28"

wb = openpyxl.load_workbook("docs/WorkspaceClone_Tracking.xlsx")


def append(sheet, rows):
    ws = wb[sheet]
    for row in rows:
        assert len(row) == ws.max_column, f"{len(row)} cells for a {ws.max_column}-column sheet: {row[0]}"
        ws.append(row)
    return len(rows)


def log(log_id, kind, component, item, message, details, tags):
    return [log_id, NOW, kind, component, item, message, details, "Assistant", tags]


append(
    "Runs",
    [
        [
            "RUN-028",
            "Build",
            "Build / Release",
            "npx tauri build (frontend + Rust release + DMG)",
            "Passed",
            f"{TODAY} 21:22",
            f"{TODAY} 21:26",
            240,
            0,
            "Workspace Clone_0.1.0_aarch64.dmg, 8.2 MB, plus the .app bundle. Ad-hoc signed "
            "(signingIdentity '-'). beforeBuildCommand ran npm run build first. One deprecation "
            "warning from tauri-build about STATIC_VCRUNTIME (release-profile only, Windows CRT "
            "config, not from any crate in the workspace); see LOG-086.",
            "",
            "src-tauri/target/release/bundle/dmg/Workspace Clone_0.1.0_aarch64.dmg",
            "Manual",
            "The debug gate the guide mandates (cargo build --workspace) remains warning-free; "
            "the release warning is toolchain-level and logged rather than patched blind.",
        ]
    ],
)

append(
    "Logs",
    [
        log(
            "LOG-086",
            "BUILD",
            "Build / Release",
            "RUN-028,FEAT-073",
            "DMG built and verified: Workspace Clone_0.1.0_aarch64.dmg, 8.2 MB",
            "npx tauri build ran the frontend build, the Rust release build (no code warnings), and "
            "bundle_dmg.sh. The only output was a tauri-build deprecation notice -- STATIC_VCRUNTIME "
            "is deprecated, use build.windows.staticVCRuntime in tauri.conf.json -- which fires in the "
            "release profile only, concerns Windows CRT linking, and originates in the tauri-build "
            "dependency rather than any crate in this workspace. Left as a documented note rather than "
            "patched blind: the config key it asks for changes the Windows installer's runtime linking, "
            "which can only be verified on the Windows laptop. The debug build gate stays warning-free.",
            "build,release,dmg,deferred-note",
        )
    ],
)

wb.save("docs/WorkspaceClone_Tracking.xlsx")
print("saved: docs/WorkspaceClone_Tracking.xlsx")
print("Runs:", wb["Runs"].max_row - 1, "Logs:", wb["Logs"].max_row - 1)